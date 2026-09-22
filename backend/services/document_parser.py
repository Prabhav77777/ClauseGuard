"""Document text extraction with page awareness.

Parses PDF and DOCX files into page-indexed text chunks. Parsing is done
synchronously in a thread pool to avoid blocking the async event loop.

Efficiency: Each document is parsed exactly once per session. The resulting
PageText list is cached in the session store for all downstream operations.
"""

import io
import logging
from fastapi import HTTPException, status
from starlette.concurrency import run_in_threadpool

from backend.models.schemas import PageText
from backend.security.validation import validate_page_count, validate_character_count

logger = logging.getLogger(__name__)


def _parse_pdf_sync(file_bytes: bytes) -> list[PageText]:
    """Extract text from PDF with page numbers. Runs in thread pool.
    
    Uses PyMuPDF for fast, accurate text extraction. Metadata is not
    included in the output — only raw text per page.
    """
    import pymupdf  # Import inside function to keep module import lightweight
    
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to parse PDF: file may be corrupted or password-protected"
        )
    
    try:
        pages = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()
            if text:  # Skip completely empty pages
                pages.append(PageText(page_number=page_num + 1, text=text))
        return pages
    finally:
        doc.close()


def _parse_docx_sync(file_bytes: bytes) -> list[PageText]:
    """Extract text from DOCX with estimated page numbers.
    
    DOCX files don't have native page boundaries, so we estimate pages
    using a character-count heuristic (~3000 chars per page). This provides
    approximate page references for citation purposes.
    """
    import docx
    
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to parse DOCX: file may be corrupted"
        )
    
    # Collect all paragraph text
    all_paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    all_paragraphs.append(text)
    
    if not all_paragraphs:
        return []
    
    # Estimate page boundaries (~3000 chars per page heuristic)
    # Efficiency: Single pass through paragraphs with running char count
    CHARS_PER_PAGE = 3000
    pages = []
    current_page_text = []
    current_char_count = 0
    current_page_num = 1
    
    for para in all_paragraphs:
        current_page_text.append(para)
        current_char_count += len(para)
        
        if current_char_count >= CHARS_PER_PAGE:
            pages.append(PageText(
                page_number=current_page_num,
                text="\n\n".join(current_page_text)
            ))
            current_page_text = []
            current_char_count = 0
            current_page_num += 1
    
    # Don't forget remaining text
    if current_page_text:
        pages.append(PageText(
            page_number=current_page_num,
            text="\n\n".join(current_page_text)
        ))
    
    return pages


async def parse_document(file_bytes: bytes, file_type: str) -> list[PageText]:
    """Parse a document into page-indexed text chunks.
    
    Runs CPU-bound parsing in a thread pool to avoid blocking the event loop.
    Validates page count and character count after parsing.
    
    Args:
        file_bytes: Raw file content
        file_type: 'pdf' or 'docx' (from magic byte detection)
    
    Returns:
        List of PageText objects with page numbers and text
    """
    if file_type == "pdf":
        # Efficiency: run_in_threadpool offloads CPU-bound PDF parsing to worker thread
        pages = await run_in_threadpool(_parse_pdf_sync, file_bytes)
    elif file_type == "docx":
        pages = await run_in_threadpool(_parse_docx_sync, file_bytes)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type: {file_type}"
        )
    
    if not pages:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No text could be extracted from the document. It may be empty, image-only, or corrupted."
        )
    
    # Post-parse validation
    validate_page_count(len(pages))
    total_text = "\n".join(p.text for p in pages)
    validate_character_count(total_text)
    
    logger.info(f"Parsed document: {len(pages)} pages, {len(total_text)} characters")
    return pages
